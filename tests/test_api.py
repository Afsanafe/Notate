import io
import json
import sqlite3
import threading
import zipfile

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader, PdfWriter
import pytest

from notate.app import create_app


def new_note(client, **overrides):
    response = client.post('/api/notes', json={'title':'Research question','content':'<p>A good question starts a discovery.</p>',**overrides})
    assert response.status_code == 201
    return response.json()


def text_annotation(client, item_id, **overrides):
    response = client.post(f'/api/items/{item_id}/annotations', json={
        'kind':'comment','color':'yellow','quote':'A good question',
        'anchor':{'type':'text','start':0,'end':15,'prefix':'','suffix':' starts'},
        'comment':'Why does this matter?',**overrides})
    assert response.status_code == 201
    return response.json()


def test_note_conflicts_do_not_overwrite_and_persist(client, app):
    item = new_note(client)
    saved = client.patch(f'/api/items/{item["id"]}',json={'revision':1,'content':'<p>The first window saved this.</p>'})
    assert saved.status_code == 200
    assert saved.json()['revision'] == 2
    stale = client.patch(f'/api/items/{item["id"]}',json={'revision':1,'content':'<p>A stale draft.</p>'})
    assert stale.status_code == 409
    with TestClient(create_app(app.state.store.directory)) as restarted:
        actual = restarted.get(f'/api/items/{item["id"]}').json()
        assert actual['content'] == '<p>The first window saved this.</p>'
        assert actual['revision'] == 2


def test_markdown_notes_preserve_source_render_tables_sync_and_conflicts(client, app):
    source = '# Research log\n\n**Evidence** and *ideas*, ~~old~~ and `code`.\n\n###### Small heading\n\n| Topic | Finding |\n| --- | --- |\n| Light | **Blue** |\n\n```python\nprint("hi")\n```\n'
    response = client.post('/api/notes',json={'title':'Markdown note','markdown':source})
    assert response.status_code == 201
    note = response.json()
    assert note['markdown'] == source
    for element in ('<h1>Research log</h1>', '<h6>Small heading</h6>', '<strong>Evidence</strong>', '<em>ideas</em>', '<s>old</s>', '<table>', '<strong>Blue</strong>', '<pre><code>'):
        assert element in note['content']
    with client.websocket_connect('/api/events') as socket:
        socket.receive_json()
        updated = client.patch(f'/api/items/{note["id"]}',json={'revision':1,'markdown':source+'\nA finding.\n'},headers={'X-Notate-Client':'markdown-editor'})
        assert updated.status_code == 200
        event = socket.receive_json()
        assert event['scope'] == 'item' and event['id'] == note['id']
    assert client.patch(f'/api/items/{note["id"]}',json={'revision':1,'markdown':'Stale source'}).status_code == 409
    with TestClient(create_app(app.state.store.directory)) as restarted:
        saved = restarted.get(f'/api/items/{note["id"]}').json()
        assert saved['markdown'] == source+'\nA finding.\n' and saved['revision'] == 2
        assert restarted.get(f'/api/items/{note["id"]}/export').json()['document']['markdown'] == saved['markdown']
        empty = restarted.patch(f'/api/items/{note["id"]}',json={'revision':2,'markdown':''}).json()
        assert empty['markdown'] == '' and empty['content'] == ''


def test_legacy_note_migration_preserves_content_revisions_and_annotations(client, app):
    original = '<h1>Research</h1><p>A good question starts a discovery. <strong>Evidence</strong> helps.</p><p><u>Underline</u> and <s>Old idea</s>.</p><table><tr><th>Key</th><th>Value</th></tr><tr><td>A</td><td>B</td></tr></table>'
    note = new_note(client,content=original)
    original = note['content']
    annotation = text_annotation(client,note['id'])
    with app.state.store.connect() as db:
        db.execute('ALTER TABLE items DROP COLUMN markdown')
    for _ in range(2):
        with TestClient(create_app(app.state.store.directory)) as restarted:
            saved = restarted.get(f'/api/items/{note["id"]}').json()
            assert saved['content'] == original
            assert saved['revision'] == 1 and saved['updated_at'] == note['updated_at']
            assert '# Research' in saved['markdown'] and '**Evidence**' in saved['markdown']
            assert '<u>Underline</u>' in saved['markdown'] and '~~Old idea~~' in saved['markdown']
            assert '| Key | Value |' in saved['markdown']
            assert restarted.get(f'/api/items/{note["id"]}/annotations').json() == [annotation]
    with TestClient(create_app(app.state.store.directory)) as restarted:
        edited = restarted.patch(f'/api/items/{note["id"]}',json={'revision':1,'markdown':saved['markdown']+'\n\nNew finding.'}).json()
        assert '<u>Underline</u>' in edited['content'] and '<s>Old idea</s>' in edited['content']
        assert 'A good question starts a discovery.' in edited['content']
        assert restarted.get(f'/api/items/{note["id"]}/annotations').json()[0]['anchor'] == annotation['anchor']


def test_markdown_note_validation_safe_rendering_and_old_draft_conversion(client, pdf_bytes):
    note = client.post('/api/notes',json={'title':'Safe Markdown','markdown':'# Note\n\n<script>alert(1)</script><img src=x onerror=alert(1)>[unsafe](javascript:alert(1))'}).json()
    assert '<script' not in note['content'] and '<img' not in note['content'] and 'javascript:' not in note['content']
    assert '<script' in note['markdown']  # Keep source verbatim, sanitize rendering.
    assert client.post('/api/notes',json={'markdown':None}).status_code == 422
    assert client.post('/api/notes',json={'markdown':'source','content':'<p>HTML</p>'}).status_code == 422
    for fields in ({'markdown':None}, {'markdown':'source','content':'<p>HTML</p>'}):
        assert client.patch(f'/api/items/{note["id"]}',json={'revision':1,**fields}).status_code == 422
    pdf = client.post('/api/import',files={'file':('test.pdf',pdf_bytes,'application/pdf')}).json()
    assert client.patch(f'/api/items/{pdf["id"]}',json={'revision':1,'markdown':'Overwrite PDF'}).status_code == 422
    converted = client.post('/api/notes/convert',json={'content':'<h2>Unsaved draft</h2><p>A <strong>finding</strong>.</p>'}).json()
    assert converted['markdown'] == '## Unsaved draft\n\nA **finding**.'
    assert client.get(f'/api/items/{note["id"]}').json()['revision'] == 1


@pytest.mark.parametrize('label', ['Answ', 'Answer'])
def test_inline_answer_markdown_is_preserved_in_notes_comments_imports_and_exports(client, app, label):
    source = f'**{label}:** *Evidence explains this.* and *ordinary italics*.'
    rendered = client.post('/api/markdown/preview',json={'body':source}).json()['html']
    assert f'<strong class="md-answer-label">{label}:</strong>' in rendered
    assert '<em class="md-answer-text">Evidence explains this.</em>' in rendered
    assert '<em>ordinary italics</em>' in rendered
    note = client.post('/api/notes',json={'title':'Inline answers','markdown':source}).json()
    assert note['markdown'] == source and note['content'] == rendered
    thread = text_annotation(client,note['id'],quote='Evidence explains this.',comment=source)
    assert thread['comments'][0]['body'] == source and thread['comments'][0]['html'] == rendered
    imported = client.post('/api/import',files={'file':('answer.md',source.encode(),'text/markdown')}).json()
    assert imported['content'] == rendered
    with TestClient(create_app(app.state.store.directory)) as restarted:
        exported = restarted.get(f'/api/items/{note["id"]}/export').json()
        assert exported['document']['markdown'] == source and exported['document']['content'] == rendered
        assert exported['annotations'][0]['comments'][0]['html'] == rendered


def test_inline_answer_styling_is_scoped_and_preserves_html_sanitization(client):
    for source in ('**Label:** *ordinary*', '`**Answ:** *code*`', '**Answ:** extra words *ordinary*', '**Answ:**\n\n*another paragraph*'):
        html = client.post('/api/markdown/preview',json={'body':source}).json()['html']
        assert 'md-answer-label' not in html and 'md-answer-text' not in html
    source = '**Answ:** *A [safe link](https://example.com).*\n\n<em class="arbitrary" style="color:red" onclick="alert(1)">unsafe</em><script>alert(1)</script>'
    html = client.post('/api/markdown/preview',json={'body':source}).json()['html']
    assert 'md-answer-text' in html and 'href="https://example.com"' in html
    assert 'arbitrary' not in html and 'style=' not in html and 'onclick=' not in html and '<script' not in html


def test_question_example_and_nested_bullets_render_and_persist(client):
    source = '**Question:** *Why does this work?*\n\n**For Example:**\n\t- *A finding.*\n\t\t- *More evidence.*\n\n- Parent\n\t- Child\n\t\t- Grandchild'
    note = client.post('/api/notes',json={'title':'Questions and examples','markdown':source}).json()
    assert note['markdown'] == source
    html = note['content']
    assert '<strong class="md-question-label">Question:</strong>' in html
    assert '<em class="md-question-text">Why does this work?</em>' in html
    assert '<strong class="md-example-label">For Example:</strong>' in html
    assert '<em class="md-example-text">A finding.</em>' in html
    assert '<ul class="md-list-circle md-example-list">' in html
    assert '<ul class="md-list-disc">' in html and '<ul class="md-list-square">' in html
    comment = text_annotation(client,note['id'],comment=source)['comments'][0]
    assert comment['body'] == source and comment['html'] == html
    imported = client.post('/api/import',files={'file':('templates.md',source.encode(),'text/markdown')}).json()
    assert imported['content'] == html
    literal = client.post('/api/markdown/preview',json={'body':'```\n'+source+'\n```'}).json()['html']
    assert 'md-question-text' not in literal and 'md-example-text' not in literal
    unsafe = client.post('/api/markdown/preview',json={'body':'<em class="md-question-text arbitrary" style="color:red">Unsafe</em>'}).json()['html']
    assert 'arbitrary' not in unsafe and 'style=' not in unsafe


def test_notepad_text_colors_round_trip_safely_with_markdown(client, app):
    colors = ['red','orange','yellow','light-green','dark-green','light-blue','dark-blue','purple']
    source = '\n\n'.join(f'<u><span class="md-text-{color}">**{color}** discovery</span></u>' for color in colors)
    note = new_note(client)
    path = f'/api/items/{note["id"]}/notepad'
    saved = client.patch(path,json={'revision':0,'markdown':source}).json()
    with TestClient(create_app(app.state.store.directory)) as restarted:
        assert restarted.get(path).json()['markdown'] == source
        assert restarted.get(f'/api/items/{note["id"]}/export').json()['notepad'] == saved
    rendered = client.post('/api/markdown/preview',json={'body':source}).json()['html']
    for color in colors:
        assert f'<u><span class="md-text-{color}"><strong>{color}</strong> discovery</span></u>' in rendered
    converted = client.post('/api/notes/convert',json={'content':rendered})
    assert converted.status_code == 200
    for color in colors:
        assert f'<span class="md-text-{color}">' in converted.json()['markdown']
    unsafe = client.post('/api/markdown/preview',json={'body':'<span class="md-text-red arbitrary" style="color:red" onclick="alert(1)">Test</span><script>alert(1)</script>'}).json()['html']
    assert 'arbitrary' not in unsafe and 'style=' not in unsafe and 'onclick=' not in unsafe and '<script' not in unsafe
    literal = client.post('/api/markdown/preview',json={'body':'```\n'+source+'\n```'}).json()['html']
    assert '<span class=' not in literal and '&lt;span' in literal


@pytest.mark.parametrize('document_kind', ['note', 'pdf'])
def test_document_notepad_is_independent_persistent_and_synced(client, app, pdf_bytes, document_kind):
    document = new_note(client) if document_kind == 'note' else client.post('/api/import',files={'file':('reading.pdf',pdf_bytes,'application/pdf')}).json()
    annotation = text_annotation(client,document['id'],**({'anchor':{'type':'pdf','pages':[{'page':1,'rects':[{'x':.1,'y':.1,'width':.3,'height':.02}]}]}} if document_kind == 'pdf' else {}))
    path = f'/api/items/{document["id"]}/notepad'
    assert client.get(path).json() == {'item_id':document['id'],'markdown':'','revision':0,'updated_at':None}
    source = '# Findings\n\n**Evidence** and *ideas*.\n\n| Topic | Finding |\n| --- | --- |\n| A | B |\n'
    with client.websocket_connect('/api/events') as socket:
        socket.receive_json()
        saved = client.patch(path,json={'revision':0,'markdown':source},headers={'X-Notate-Client':'pad-window'})
        assert saved.status_code == 200 and saved.json()['markdown'] == source and saved.json()['revision'] == 1
        event = socket.receive_json()
        assert event['scope'] == 'notepad' and event['id'] == document['id'] and event['client'] == 'pad-window'
    assert client.patch(path,json={'revision':0,'markdown':'Stale notepad'}).status_code == 409
    assert client.get(f'/api/items/{document["id"]}').json() == document
    assert client.get(f'/api/items/{document["id"]}/annotations').json() == [annotation]
    with TestClient(create_app(app.state.store.directory)) as restarted:
        assert restarted.get(path).json()['markdown'] == source
        exported = restarted.get(f'/api/items/{document["id"]}/export').json()
        assert exported['notepad']['markdown'] == source and exported['annotations'] == [annotation]
        assert restarted.patch(path,json={'revision':1,'markdown':''}).json()['revision'] == 2
    assert client.patch(path,json={'revision':2,'markdown':None}).status_code == 422
    assert client.patch('/api/items/missing/notepad',json={'revision':0,'markdown':'Draft'}).status_code == 404


def test_notepads_migrate_isolate_restore_backups_and_cascade(client, app, tmp_path):
    first, second = new_note(client), new_note(client,title='Other document')
    with app.state.store.connect() as db:
        db.execute('DROP TABLE notepads')
    with TestClient(create_app(app.state.store.directory)) as migrated:
        for note, text in ((first,'**First** findings.'),(second,'# Second findings')):
            response = migrated.patch(f'/api/items/{note["id"]}/notepad',json={'revision':0,'markdown':text})
            assert response.status_code == 200
        backup = migrated.get('/api/backup')
        with zipfile.ZipFile(io.BytesIO(backup.content)) as archive:
            archive.extractall(tmp_path/'restored-pads')
        assert migrated.delete(f'/api/items/{first["id"]}').status_code == 200
        with app.state.store.connect() as db:
            assert [row['item_id'] for row in db.execute('SELECT item_id FROM notepads')] == [second['id']]
    with TestClient(create_app(tmp_path/'restored-pads')) as restored:
        assert restored.get(f'/api/items/{first["id"]}/notepad').json()['markdown'] == '**First** findings.'
        assert restored.get(f'/api/items/{second["id"]}/notepad').json()['markdown'] == '# Second findings'


def test_two_initial_notepad_saves_cannot_overwrite_one_another(client):
    from concurrent.futures import ThreadPoolExecutor
    note = new_note(client)
    path = f'/api/items/{note["id"]}/notepad'
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(client.patch,path,json={'revision':0,'markdown':source}) for source in ('First draft','Second draft')]
        responses = [future.result() for future in futures]
    assert sorted(response.status_code for response in responses) == [200,409]
    winning = next(response.json() for response in responses if response.status_code == 200)
    assert client.get(path).json() == winning


@pytest.mark.parametrize('document_kind', ['note', 'text', 'pdf'])
def test_passage_pad_links_commit_both_sides_sync_and_restore(client, app, tmp_path, pdf_bytes, document_kind):
    if document_kind == 'note':
        document = new_note(client)
    else:
        filename, data, mime = ('reading.pdf',pdf_bytes,'application/pdf') if document_kind == 'pdf' else ('reading.txt',b'A good question starts a discovery.','text/plain')
        document = client.post('/api/import',files={'file':(filename,data,mime)}).json()
    item_id = document['id']
    anchor = {'type':'pdf','pages':[{'page':1,'rects':[{'x':.1,'y':.1,'width':.3,'height':.02}]}]} if document_kind == 'pdf' else {'type':'text','start':0,'end':15}
    with client.websocket_connect('/api/events') as socket:
        socket.receive_json()
        response = client.post(f'/api/items/{item_id}/pad-links',json={'quote':'A good question','anchor':anchor,'revision':0},headers={'X-Notate-Client':'link-window'})
        assert response.status_code == 201
        result = response.json(); annotation, pad = result['annotation'], result['notepad']
        events = [socket.receive_json(),socket.receive_json()]
        assert [event['scope'] for event in events] == ['notepad','annotations']
        assert all(event['id'] == item_id and event['client'] == 'link-window' for event in events)
    assert annotation['kind'] == 'pad' and annotation['color'] == 'orange'
    assert annotation['anchor'] == anchor and annotation['quote'] == 'A good question'
    assert pad['markdown'] == f'[A good question](#notate-passage-{annotation["id"]})\n\n'
    assert pad['revision'] == 1
    # Metadata lives in the ordinary workspace backup and JSON export.
    exported = client.get(f'/api/items/{item_id}/export').json()
    assert exported['notepad'] == pad and exported['annotations'] == [annotation]
    backup = client.get('/api/backup')
    with zipfile.ZipFile(io.BytesIO(backup.content)) as archive:
        archive.extractall(tmp_path/'linked-workspace')
    with TestClient(create_app(tmp_path/'linked-workspace')) as restored:
        assert restored.get(f'/api/items/{item_id}/notepad').json() == pad
        assert restored.get(f'/api/items/{item_id}/annotations').json() == [annotation]
    assert client.delete(f'/api/items/{item_id}').status_code == 200
    with app.state.store.connect() as db:
        assert db.execute('SELECT COUNT(*) FROM annotations WHERE item_id=?',(item_id,)).fetchone()[0] == 0
        assert db.execute('SELECT COUNT(*) FROM notepads WHERE item_id=?',(item_id,)).fetchone()[0] == 0


def test_link_existing_pad_words_utf16_edit_and_unlink_preserves_text(client):
    document = new_note(client); item_id = document['id']; path = f'/api/items/{item_id}/notepad'
    source = '# Findings\n\n🌟 A hypothesis and an existing [site](https://example.test).\n'
    client.patch(path,json={'revision':0,'markdown':source})
    start = len(source[:source.index('hypothesis')].encode('utf-16-le'))//2
    body = {'quote':'A good question','anchor':{'type':'text','start':0,'end':15},'revision':1,
            'selection':{'start':start,'end':start+10,'text':'hypothesis'}}
    response = client.post(f'/api/items/{item_id}/pad-links',json=body)
    assert response.status_code == 201
    result = response.json(); link_id = result['annotation']['id']
    assert result['notepad']['markdown'] == source.replace('hypothesis',f'[hypothesis](#notate-passage-{link_id})')
    # Editing or moving the linked label retains its destination in Markdown.
    edited = result['notepad']['markdown'].replace('hypothesis','[nested] idea')
    assert client.patch(path,json={'revision':2,'markdown':edited}).status_code == 200
    assert client.delete(f'/api/annotations/{link_id}').status_code == 200
    assert client.get(path).json()['markdown'] == source.replace('hypothesis','[nested] idea')
    assert client.get(path).json()['revision'] == 4
    assert client.get(f'/api/items/{item_id}/annotations').json() == []


def test_pad_link_validation_and_conflicts_leave_no_partial_links(client):
    from concurrent.futures import ThreadPoolExecutor
    document = new_note(client); item_id = document['id']; path = f'/api/items/{item_id}/pad-links'
    body = {'quote':'A good question','anchor':{'type':'text','start':0,'end':15},'revision':0}
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda _: client.post(path,json=body),range(2)))
    assert sorted(response.status_code for response in responses) == [201,409]
    winner = next(response.json() for response in responses if response.status_code == 201)
    assert client.get(f'/api/items/{item_id}/annotations').json() == [winner['annotation']]
    assert client.get(f'/api/items/{item_id}/notepad').json() == winner['notepad']
    invalid = {**body,'revision':1,'selection':{'start':999999,'end':1000000,'text':'x'}}
    assert client.post(path,json=invalid).status_code == 422
    assert client.post(path,json={**invalid,'selection':{'start':0,'end':1,'text':'wrong'}}).status_code == 422
    assert client.post(path,json={**body,'revision':1,'quote':'   '}).status_code == 422
    assert client.get(f'/api/items/{item_id}/notepad').json() == winner['notepad']
    assert client.get(f'/api/items/{item_id}/annotations').json() == [winner['annotation']]


def test_pad_link_response_survives_document_deletion_after_commit(client, app, monkeypatch):
    document = new_note(client); item_id = document['id']
    publish = app.state.hub.publish

    async def delete_after_commit(request, scope, changed_id):
        if scope == 'annotations' and changed_id == item_id:
            with app.state.store.connect() as db:
                db.execute('DELETE FROM items WHERE id=?',(item_id,))
        await publish(request,scope,changed_id)

    monkeypatch.setattr(app.state.hub,'publish',delete_after_commit)
    response = client.post(f'/api/items/{item_id}/pad-links',json={
        'quote':'A good question','anchor':{'type':'text','start':0,'end':15},'revision':0})
    assert response.status_code == 201
    assert response.json()['annotation']['kind'] == 'pad'
    assert client.get(f'/api/items/{item_id}/notepad').status_code == 404


def test_nested_folder_cycles_and_nonempty_delete(client):
    root = client.post('/api/folders',json={'name':'Research'}).json()
    child = client.post('/api/folders',json={'name':'Questions','parent_id':root['id']}).json()
    cycle = client.patch(f'/api/folders/{root["id"]}',json={'name':'Research','parent_id':child['id']})
    assert cycle.status_code == 409
    note = new_note(client,folder_id=child['id'])
    assert client.delete(f'/api/folders/{child["id"]}').status_code == 409
    assert client.patch(f'/api/items/{note["id"]}',json={'revision':1,'folder_id':None}).status_code == 200
    assert client.delete(f'/api/folders/{child["id"]}').status_code == 200
    assert client.delete(f'/api/folders/{root["id"]}').status_code == 200
    assert client.get(f'/api/items/{note["id"]}').json()['folder_id'] is None


def test_threads_replies_resolving_and_cascade(client, app):
    note = new_note(client)
    annotation = text_annotation(client,note['id'])
    response = client.post(f'/api/annotations/{annotation["id"]}/comments',json={'body':'It helps connect observations.'})
    assert response.status_code == 201
    assert client.patch(f'/api/annotations/{annotation["id"]}',json={'resolved':True}).status_code == 200
    records = client.get(f'/api/items/{note["id"]}/annotations').json()
    assert records[0]['resolved'] is True
    assert [c['body'] for c in records[0]['comments']] == ['Why does this matter?','It helps connect observations.']
    assert client.delete(f'/api/items/{note["id"]}').status_code == 200
    with app.state.store.connect() as db:
        assert db.execute('SELECT COUNT(*) FROM annotations').fetchone()[0] == 0
        assert db.execute('SELECT COUNT(*) FROM comments').fetchone()[0] == 0


def test_answer_type_markdown_changes_persist_sync_and_export(client, app):
    note = new_note(client)
    annotation = text_annotation(client,note['id'],comment='**Question:** why?')
    first = annotation['comments'][0]
    assert first['kind'] == 'reply' and '<strong>Question:</strong>' in first['html']
    body = '**Finding:** evidence helps.\n\n- Observe\n- Compare\n\n`result = 2`'
    answer = client.post(f'/api/annotations/{annotation["id"]}/comments',json={'body':body,'kind':'answer'}).json()
    assert answer['kind'] == 'answer' and answer['body'] == body
    assert '<strong>Finding:</strong>' in answer['html'] and '<ul class="md-list-disc">' in answer['html'] and '<code>result = 2</code>' in answer['html']
    with client.websocket_connect('/api/events') as socket:
        socket.receive_json()
        changed = client.patch(f'/api/comments/{answer["id"]}',json={'kind':'reply'},headers={'X-Notate-Client':'other-window'})
        assert changed.status_code == 200 and changed.json()['body'] == body and changed.json()['kind'] == 'reply'
        event = socket.receive_json()
        assert event['scope'] == 'annotations' and event['id'] == note['id'] and event['client'] == 'other-window'
    assert client.patch(f'/api/comments/{answer["id"]}',json={'kind':'answer'}).status_code == 200
    with TestClient(create_app(app.state.store.directory)) as restarted:
        exported = restarted.get(f'/api/items/{note["id"]}/export').json()['annotations'][0]['comments']
        assert exported[1]['kind'] == 'answer' and exported[1]['body'] == body
    assert client.patch(f'/api/comments/{answer["id"]}',json={'kind':'unexpected'}).status_code == 422
    assert client.patch('/api/comments/missing',json={'kind':'answer'}).status_code == 404


def test_existing_comment_database_migrates_without_losing_threads(client, app):
    note = new_note(client)
    annotation = text_annotation(client,note['id'],comment='**Original** question.')
    # Recreate the schema used by earlier Notate installations, preserving its rows.
    with app.state.store.connect() as db:
        db.executescript('''
            CREATE TABLE legacy_comments (id TEXT PRIMARY KEY,
                annotation_id TEXT NOT NULL REFERENCES annotations(id) ON DELETE CASCADE,
                body TEXT NOT NULL, created_at TEXT NOT NULL);
            INSERT INTO legacy_comments SELECT id,annotation_id,body,created_at FROM comments;
            DROP TABLE comments;
            ALTER TABLE legacy_comments RENAME TO comments;
        ''')
    for _ in range(2):
        with TestClient(create_app(app.state.store.directory)) as restarted:
            comments = restarted.get(f'/api/items/{note["id"]}/annotations').json()[0]['comments']
            assert len(comments) == 1
            assert comments[0]['id'] == annotation['comments'][0]['id']
            assert comments[0]['body'] == '**Original** question.' and comments[0]['kind'] == 'reply'
            assert '<strong>Original</strong>' in comments[0]['html']
    with TestClient(create_app(app.state.store.directory)) as restarted:
        assert restarted.patch(f'/api/comments/{comments[0]["id"]}',json={'kind':'answer'}).status_code == 200


def test_markdown_preview_matches_saved_rendering_and_sanitizes_html(client):
    body = '# Heading\n\n**bold** and *italic* and ~~old~~\nnext line\n\n> a quotation\n\n1. first\n2. second\n\n```python\nprint("hello")\n```\n\n| Key | Value |\n| --- | --- |\n| A | B |\n\n[safe](https://example.com)'
    rendered = client.post('/api/markdown/preview',json={'body':body})
    assert rendered.status_code == 200
    html = rendered.json()['html']
    for tag in ('<h1>', '<strong>', '<em>', '<s>', '<br>', '<blockquote>', '<ol>', '<pre><code>', '<table>'):
        assert tag in html
    assert 'href="https://example.com"' in html
    note = new_note(client)
    annotation = text_annotation(client,note['id'],comment=body,comment_kind='answer')
    assert annotation['comments'][0]['html'] == html and annotation['comments'][0]['kind'] == 'answer'
    unsafe = '<script>alert(1)</script><img src=x onerror=alert(1)>[unsafe](javascript:alert(1))'
    sanitized = client.post('/api/markdown/preview',json={'body':unsafe}).json()['html']
    assert '<script' not in sanitized and '<img' not in sanitized and 'onerror' not in sanitized and 'javascript:' not in sanitized
    indented = client.post(f'/api/annotations/{annotation["id"]}/comments',json={'body':'    indented code'}).json()
    assert indented['body'] == '    indented code' and '<pre><code>indented code' in indented['html']
    assert client.post('/api/markdown/preview',json={'body':'x'*20001}).status_code == 422


def test_pdf_bytes_preserved_range_requests_and_anchors(client, pdf_bytes):
    response = client.post('/api/import',files={'file':('field-guide.pdf',pdf_bytes,'application/pdf')})
    assert response.status_code == 201
    item = response.json()
    assert item['has_pdf'] is True and item['page_count'] == 2
    assert 'file_path' not in item
    assert client.get(f'/api/items/{item["id"]}/file').content == pdf_bytes
    partial = client.get(f'/api/items/{item["id"]}/file',headers={'Range':'bytes=0-99'})
    assert partial.status_code == 206 and partial.content == pdf_bytes[:100]
    anchor = {'type':'pdf','pages':[{'page':1,'rects':[{'x':.1,'y':.1,'width':.3,'height':.02}]}]}
    annotation = text_annotation(client,item['id'],anchor=anchor,quote='Reading is a conversation',kind='underline',comment=None)
    assert annotation['anchor'] == anchor
    anchor['pages'][0]['page'] = 3
    assert client.post(f'/api/items/{item["id"]}/annotations',json={'kind':'highlight','quote':'text','anchor':anchor}).status_code == 422
    assert client.patch(f'/api/items/{item["id"]}',json={'revision':1,'content':'Replaced'}).status_code == 422


@pytest.mark.parametrize('algorithm', ['AES-128', 'AES-256'])
def test_openable_aes_pdf_imports_without_changing_original(client, pdf_bytes, algorithm):
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf_bytes)))
    writer.encrypt(user_password='', owner_password='document-owner', algorithm=algorithm)
    output = io.BytesIO()
    writer.write(output)
    encrypted = output.getvalue()
    response = client.post('/api/import',files={'file':('aes-encrypted.pdf',encrypted,'application/pdf')})
    assert response.status_code == 201, response.json()
    item = response.json()
    assert item['has_pdf'] is True and item['page_count'] == 2
    assert client.get(f'/api/items/{item["id"]}/file?download=true').content == encrypted


def test_aes_pdf_with_opening_password_still_requires_unlocked_copy(client, app, pdf_bytes):
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf_bytes)))
    writer.encrypt(user_password='opening-password', owner_password='document-owner', algorithm='AES-256')
    output = io.BytesIO()
    writer.write(output)
    response = client.post('/api/import',files={'file':('password-protected.pdf',output.getvalue(),'application/pdf')})
    assert response.status_code == 422
    assert 'password protected' in response.json()['detail']
    assert list(app.state.store.files.iterdir()) == []


def test_missing_aes_dependency_has_actionable_error(client, monkeypatch, pdf_bytes):
    from pypdf.errors import DependencyError

    def missing_crypto(path):
        raise DependencyError('cryptography is required for AES algorithm')

    monkeypatch.setattr('notate.documents.PdfReader',missing_crypto)
    response = client.post('/api/import',files={'file':('aes.pdf',pdf_bytes,'application/pdf')})
    assert response.status_code == 422
    assert 'AES support is missing' in response.json()['detail']
    assert 'restart' in response.json()['detail']


def test_docx_fallback_preserves_text_formatting_tables(client, monkeypatch):
    monkeypatch.setattr('notate.documents.convert_office',lambda *args:None)
    document = Document()
    document.add_heading('A Word document',0)
    paragraph = document.add_paragraph('A useful ')
    paragraph.add_run('finding').bold = True
    table = document.add_table(rows=1,cols=2)
    table.cell(0,0).text = 'Question'
    table.cell(0,1).text = 'Answer'
    data = io.BytesIO()
    document.save(data)
    imported = client.post('/api/import',files={'file':('research.docx',data.getvalue())})
    assert imported.status_code == 201
    item = client.get(f'/api/items/{imported.json()["id"]}').json()
    assert not item['has_pdf']
    assert '<h1>A Word document</h1>' in item['content']
    assert '<strong>finding</strong>' in item['content']
    assert '<table>' in item['content'] and 'Answer' in item['content']


def test_bad_imports_are_removed(client, app):
    assert client.post('/api/import',files={'file':('broken.pdf',b'not a PDF')}).status_code == 422
    assert client.post('/api/import',files={'file':('broken.docx',b'not a Word archive')}).status_code == 422
    assert client.post('/api/import',files={'file':('empty.txt',b'')}).status_code == 422
    assert client.post('/api/import',files={'file':('bad.html',b'<script>alert(1)</script>')}).status_code == 422
    assert list(app.state.store.files.iterdir()) == []


def test_local_sync_events_identify_origin_and_scope(client):
    with client.websocket_connect('/api/events') as socket:
        assert socket.receive_json()['type'] == 'connected'
        response = client.post('/api/notes',headers={'X-Notate-Client':'desktop-window'},json={'title':'Live note'})
        assert response.status_code == 201
        event = socket.receive_json()
        assert event['scope'] == 'workspace' and event['client'] == 'desktop-window'
        assert event['sequence'] == 1
        socket.send_text('ping')
        assert socket.receive_json() == {'type':'pong'}


def test_backup_is_restorable_and_exports_threads(client, tmp_path, pdf_bytes):
    note = new_note(client)
    text_annotation(client,note['id'])
    imported = client.post('/api/import',files={'file':('original.pdf',pdf_bytes)}).json()
    exported = client.get(f'/api/items/{note["id"]}/export').json()
    assert exported['annotations'][0]['comments'][0]['body'] == 'Why does this matter?'
    response = client.get('/api/backup')
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        archive.extractall(tmp_path / 'restored')
    with TestClient(create_app(tmp_path / 'restored')) as restored:
        assert restored.get(f'/api/items/{imported["id"]}/file').content == pdf_bytes
        assert len(restored.get(f'/api/items/{note["id"]}/annotations').json()[0]['comments']) == 1


def test_backup_keeps_snapshot_files_when_deleted_concurrently(client, app, pdf_bytes, monkeypatch):
    imported = client.post('/api/import',files={'file':('source.pdf',pdf_bytes)}).json()
    copying = threading.Event()
    resume = threading.Event()
    original_write = zipfile.ZipFile.write
    results = {}

    def delayed_write(archive, filename, arcname=None, *args, **kwargs):
        if arcname and arcname.endswith('.pdf'):
            copying.set()
            assert resume.wait(5)
        return original_write(archive,filename,arcname,*args,**kwargs)

    monkeypatch.setattr(zipfile.ZipFile,'write',delayed_write)
    backup = threading.Thread(target=lambda:results.update(backup=client.get('/api/backup')))
    delete = threading.Thread(target=lambda:results.update(delete=client.delete(f'/api/items/{imported["id"]}')))
    backup.start()
    try:
        assert copying.wait(5)
        delete.start()
        delete.join(.1)
        assert delete.is_alive(), 'Deletion must wait until the backup copies its source file.'
    finally:
        resume.set()
        backup.join(5)
        if delete.ident:
            delete.join(5)
    assert results['backup'].status_code == 200
    assert results['delete'].status_code == 200
    with zipfile.ZipFile(io.BytesIO(results['backup'].content)) as archive:
        pdf_name = next(name for name in archive.namelist() if name.endswith('.pdf'))
        assert archive.read(pdf_name) == pdf_bytes
    assert client.get(f'/api/items/{imported["id"]}').status_code == 404


def test_note_html_and_cross_origin_mutations_are_restricted(client):
    note = new_note(client,content='<p onclick="alert(1)">Safe</p><img src=x onerror=alert(1)><a href="javascript:alert(1)">link</a>')
    assert 'onclick' not in note['content'] and 'onerror' not in note['content'] and 'javascript:' not in note['content']
    response = client.post('/api/notes',headers={'Origin':'https://elsewhere.example'},json={'title':'Foreign request'})
    assert response.status_code == 403


def test_network_workspace_password(tmp_path):
    with TestClient(create_app(tmp_path / 'locked',token='workspace-secret')) as locked:
        assert locked.get('/').status_code == 200
        assert locked.get('/api/health').json()['locked'] is True
        assert locked.get('/api/workspace').status_code == 401
        assert locked.post('/api/session',json={'token':'wrong'}).status_code == 401
        response = locked.post('/api/session',json={'token':'workspace-secret'})
        assert response.status_code == 200
        assert 'HttpOnly' in response.headers['set-cookie']
        assert locked.get('/api/workspace').status_code == 200
        with locked.websocket_connect('/api/events') as socket:
            assert socket.receive_json()['type'] == 'connected'


def test_empty_comments_and_invalid_anchors_rejected(client):
    note = new_note(client)
    assert client.post(f'/api/items/{note["id"]}/annotations',json={'kind':'comment','quote':'A','comment':' ','anchor':{'type':'text','start':0,'end':1}}).status_code == 422
    assert client.post(f'/api/items/{note["id"]}/annotations',json={'kind':'highlight','quote':'A','anchor':{'type':'text','start':10,'end':1}}).status_code == 422
    annotation = text_annotation(client,note['id'])
    assert client.post(f'/api/annotations/{annotation["id"]}/comments',json={'body':'   '}).status_code == 422
